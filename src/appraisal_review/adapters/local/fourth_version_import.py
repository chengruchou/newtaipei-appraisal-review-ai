"""Import the fourth-version Shulin results into the snapshot the writers already read.

The three workbooks in ``verify/output/fourthVersion`` are the official templates with
results already filled in, so their visible sheets carry the template's own geometry.
This importer therefore reads them *through the existing table mappings*: for every
``CellBinding`` a mapping declares, the same cell is read from the source sheet and
becomes that binding's snapshot entry. No coordinate is invented here, and no second
Excel renderer exists - ``workbook_writer.fill_workbook`` still does the drawing.

What the import deliberately does NOT do:

- It does not confirm anything. Values arrive as ``imported_result``: traceable to a
  pinned commit, file digest, sheet and cell, and still subject to the formal approval
  gate. Nothing is recorded as ``human_confirmed``, because no human confirmed it.
- It does not turn absence into zero. A blank, an em dash, ``NA`` and friends stay
  absent and are reported as gaps with their reason, while a real ``0`` is kept as the
  number ``0`` - not as the "confirmed zero" state, which asserts a human decision.
- It does not rescale a percentage twice. The binding's own ``value_kind`` says how the
  source cell is written: a ``percentage_fraction`` cell holds 0.09 and the snapshot
  stores 9 percent points, while a ``percentage_points`` cell already holds 9 and is
  stored unchanged. The writer then renders each back in that cell's own convention.
"""

# The fullwidth punctuation below is the assignment's own: these strings are read by
# Chinese-reading reviewers and are compared against the source workbooks verbatim.
# ruff: noqa: RUF001, RUF002

from __future__ import annotations

import hashlib
import io
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

import openpyxl

from appraisal_review.domain.calculation_snapshot import SnapshotEntry
from appraisal_review.domain.official_table_mapping import CellBinding, TableMapping

#: The reviewed commit these workbooks were taken from; recorded in every entry's trace.
SOURCE_COMMIT = "a6e4dc85cb688d857dfae1ef54d234b18f0de078"

#: Text a source cell uses to say "no value", as opposed to a real zero. Kept explicit:
#: the assignment distinguishes "not surveyed" from "surveyed and found to be none",
#: and collapsing either into 0 would invent a finding.
ABSENT_TEXT = frozenset({"", "-", "--", "—", "–", "－", "na", "n/a", "nil", "無資料", "資料不足"})

#: The label the UI and the export attach to everything this importer produced.
IMPORT_LABEL = "第四版匯入結果／待審查"


@dataclass(frozen=True)
class SourceCell:
    """One cell as the source workbook actually holds it, before any conversion."""

    sheet: str
    cell: str
    raw: object
    note: str = ""


@dataclass(frozen=True)
class ImportedTable:
    """One table's entries and gaps, plus the provenance of the file they came from."""

    table: str
    section: str
    file_name: str
    file_digest: str
    sheet_name: str
    entries: dict[str, SnapshotEntry]
    gaps: dict[str, str]
    #: Bindings whose source cell held a value this importer chose not to convert.
    unmapped: dict[str, str]


def read_basis_notes(workbook: openpyxl.Workbook) -> dict[tuple[str, str], str]:
    """Index the workbook's own 填表依據 sheet by (section, cell).

    The sheet is the teammates' filling rationale, not a machine contract, so a shape
    this does not recognise yields no notes rather than an error: the note enriches an
    entry's trace and never decides whether a value is imported.
    """
    if "填表依據" not in workbook.sheetnames:
        return {}
    sheet = workbook["填表依據"]
    header_row = None
    for row in range(1, min(sheet.max_row, 40) + 1):
        if str(sheet.cell(row=row, column=2).value or "").strip() == "儲存格":
            header_row = row
            break
    if header_row is None:
        return {}
    notes: dict[tuple[str, str], str] = {}
    for row in range(header_row + 1, sheet.max_row + 1):
        section = str(sheet.cell(row=row, column=1).value or "").strip()
        cell = str(sheet.cell(row=row, column=2).value or "").strip()
        if not cell:
            continue
        field = str(sheet.cell(row=row, column=3).value or "").strip()
        source = str(sheet.cell(row=row, column=5).value or "").strip()
        basis = str(sheet.cell(row=row, column=6).value or "").strip()
        nature = str(sheet.cell(row=row, column=8).value or "").strip()
        parts = [part for part in (field, source, basis, nature) if part]
        notes[(section, cell)] = " ｜ ".join(parts)
    return notes


def _is_absent(raw: object) -> bool:
    if raw is None:
        return True
    if isinstance(raw, str):
        return raw.strip().casefold() in ABSENT_TEXT
    return False


def _as_decimal(raw: object) -> tuple[Decimal, bool] | None:
    """The number plus whether the source spelled it with a percent sign.

    A textual "50%" already states percent points; only a bare numeric in a
    fraction-convention cell still needs the points shift. Losing this distinction
    turned 建蔽率 50% into 5000% in an early draft of this importer.
    """
    if isinstance(raw, bool):
        return None
    if isinstance(raw, (int, float, Decimal)):
        return Decimal(str(raw)), False
    if isinstance(raw, str):
        stripped = raw.strip().replace(",", "")
        spelled_percent = stripped.endswith("%")
        try:
            return Decimal(stripped.rstrip("%")), spelled_percent
        except InvalidOperation:
            return None
    return None


def _convert(binding: CellBinding, raw: object) -> tuple[str | Decimal | None, str | None]:
    """The binding's declared kind decides the conversion; unit follows the binding."""
    kind = binding.value_kind
    if kind in {"text", "date"}:
        if isinstance(raw, (datetime, date)):
            return raw.strftime("%Y-%m-%d"), binding.unit
        text = str(raw).strip()
        return (text, binding.unit) if text else (None, None)
    parsed = _as_decimal(raw)
    if parsed is None:
        return None, None
    number, spelled_percent = parsed
    if kind == "percentage_fraction":
        # A bare numeric in a fraction-convention cell states a fraction, so the
        # snapshot's percent-point currency needs the shift; "50%" spelled out is
        # already points and shifts nothing.
        return (number if spelled_percent else number * Decimal(100)), binding.unit
    if kind == "percentage_points":
        # Already points - rescaling here would be the classic double conversion.
        return number, binding.unit
    if kind == "integer":
        if number != number.to_integral_value():
            return None, None
        return Decimal(number.to_integral_value()), binding.unit
    return number, binding.unit


def import_table(
    *,
    source_bytes: bytes,
    mapping: TableMapping,
    source_sheet: str,
    section: str,
    file_name: str,
    source_prefix: str | None = None,
) -> ImportedTable:
    """Read every cell the mapping binds, from the matching sheet of the source file.

    ``source_sheet`` is named explicitly because the source and the official template
    do not always agree on it - table 5's sheet is 表5區域因素明細表＿一般住宅 in the
    source and 表5-1區域因素明細表(住) in the template - so the caller states the pair
    rather than letting a lookup guess.

    ``source_prefix`` rebases each binding's snapshot key, which is how table 3 keeps
    all four sections: the same 198 bindings are read once per section sheet into
    separate keys instead of the later sections overwriting the first.
    """
    digest = hashlib.sha256(source_bytes).hexdigest()
    workbook = openpyxl.load_workbook(io.BytesIO(source_bytes), data_only=True)
    if source_sheet not in workbook.sheetnames:
        raise KeyError(f"{file_name} has no sheet {source_sheet!r}")
    sheet = workbook[source_sheet]
    notes = read_basis_notes(workbook)

    entries: dict[str, SnapshotEntry] = {}
    gaps: dict[str, str] = {}
    unmapped: dict[str, str] = {}
    for binding in mapping.bindings:
        key = binding.source
        if source_prefix is not None:
            key = _rebase(key, source_prefix)
        raw: Any = sheet[binding.cell].value
        if _is_absent(raw):
            gaps[key] = (
                f"{IMPORT_LABEL}：來源 {file_name} {source_sheet}!{binding.cell}"
                f"（{binding.label}）未填，依原依據維持缺值。"
            )
            continue
        value, unit = _convert(binding, raw)
        if value is None:
            unmapped[key] = (
                f"{source_sheet}!{binding.cell} 值 {raw!r} 不符 {binding.value_kind} 型別，未匯入。"
            )
            gaps[key] = (
                f"{IMPORT_LABEL}：來源 {file_name} {source_sheet}!{binding.cell}"
                f"（{binding.label}）值無法依 {binding.value_kind} 解讀，維持缺值。"
            )
            continue
        note = notes.get((section, binding.cell)) or notes.get(("全表", binding.cell), "")
        trace = (
            f"{IMPORT_LABEL}｜fourthVersion@{SOURCE_COMMIT[:12]}"
            f"｜{file_name}#{digest[:12]}｜{source_sheet}!{binding.cell}｜原值={raw!r}"
        )
        if note:
            trace = f"{trace}｜{note}"
        entries[key] = SnapshotEntry(
            state="present",
            value=value,
            unit=unit,
            origin="imported_result",
            trace=trace[:2048],
        )
    return ImportedTable(
        table=mapping.table,
        section=section,
        file_name=file_name,
        file_digest=digest,
        sheet_name=source_sheet,
        entries=entries,
        gaps=gaps,
        unmapped=unmapped,
    )


def _rebase(source_key: str, prefix: str) -> str:
    """``table_3.case.zoning`` under prefix ``P002`` becomes ``table_3.P002.zoning``."""
    head, _, tail = source_key.partition(".")
    _, _, field = tail.partition(".")
    return f"{head}.{prefix}.{field}" if field else f"{head}.{prefix}"


def section_mapping(mapping: TableMapping, prefix: str) -> TableMapping:
    """The same mapping addressing one section's keys, for rendering that section."""
    return mapping.model_copy(
        update={
            "bindings": tuple(
                binding.model_copy(update={"source": _rebase(binding.source, prefix)})
                for binding in mapping.bindings
            )
        }
    )
