"""The official writer fills real template copies without losing fidelity.

Every test runs against the organizer's real workbooks under
``artifacts/official-templates`` (0700, gitignored). When the directory is
absent - CI without the materials - the whole module skips cleanly.
"""

from __future__ import annotations

import io
import json
import re
import zipfile
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest

from appraisal_review.adapters.local.workbook_writer import (
    FilledWorkbook,
    WorkbookWriteError,
    fill_workbook,
)
from appraisal_review.domain.calculation_snapshot import (
    CalculationSnapshot,
    SnapshotEntry,
    SnapshotSubject,
)
from appraisal_review.domain.official_table_mapping import TableMapping
from appraisal_review.domain.service_contracts import RevisionReference

TEMPLATES = Path(__file__).resolve().parents[2] / "artifacts" / "official-templates"
MAPPINGS = Path(__file__).resolve().parents[2] / "configs" / "mappings"
PCT = "percent_points"

pytestmark = pytest.mark.skipif(
    not TEMPLATES.is_dir(), reason="official templates are not materialized"
)


def _mapping(name: str) -> TableMapping:
    return TableMapping.model_validate(json.loads((MAPPINGS / name).read_text(encoding="utf-8")))


def _template(prefix: str) -> bytes:
    return next(TEMPLATES.glob(f"{prefix}*.xlsx")).read_bytes()


def _snapshot(
    entries: dict[str, SnapshotEntry], gaps: dict[str, str] | None = None
) -> CalculationSnapshot:
    return CalculationSnapshot(
        revision=RevisionReference(case_id="case-1", revision_id="r1", material_digest="c" * 64),
        district="金山區",
        valuation_date=date(2025, 9, 1),
        rule_bundle_id="rule-bundle",
        rule_bundle_version="1",
        subjects=(
            SnapshotSubject(subject_id="P001", role="comparison_base", label="比準地"),
            SnapshotSubject(subject_id="P002", role="comparable", label="比較標的1"),
            SnapshotSubject(subject_id="P003", role="comparable", label="比較標的2"),
            SnapshotSubject(subject_id="P004", role="comparable", label="比較標的3"),
        ),
        entries=entries,
        gaps=gaps or {},
    )


def _present(value: Decimal | str, unit: str | None = None) -> SnapshotEntry:
    return SnapshotEntry(
        state="present", value=value, unit=unit, origin="computed", trace="fixture arithmetic"
    )


def _sheet_xml(result: FilledWorkbook, template: bytes) -> str:
    with (
        zipfile.ZipFile(io.BytesIO(template)) as before,
        zipfile.ZipFile(io.BytesIO(result.content)) as after,
    ):
        changed = [name for name in before.namelist() if before.read(name) != after.read(name)]
        assert len(changed) == 1, f"exactly one part may change, saw {changed}"
        assert before.namelist() == after.namelist()
        return after.read(changed[0]).decode("utf-8")


def _cell(xml: str, reference: str) -> str:
    match = re.search(rf'<c r="{reference}"[^>/]*(?:/>|>.*?</c>)', xml, flags=re.DOTALL)
    assert match is not None, f"{reference} not found"
    return match.group(0)


def test_table5_round_trip_preserves_every_untouched_part() -> None:
    template = _template("表5")
    mapping = _mapping("table5-v1.json")
    result = fill_workbook(
        template,
        mapping,
        _snapshot(
            {
                "table_5.case.case_number": _present("1140901-99-001"),
                "table_5.P001.urban_plan_grade": _present(Decimal("1")),
                "table_5.P002.urban_plan_grade_label": _present("優"),
                "table_5.P002.urban_plan_adjustment_pct": _present(Decimal("5"), unit=PCT),
            }
        ),
    )
    xml = _sheet_xml(result, template)  # asserts one changed part, identical part list
    assert _cell(xml, "B2").endswith(
        't="inlineStr"><is><t xml:space="preserve">1140901-99-001</t></is></c>'
    )
    assert "<v>1</v>" in _cell(xml, "C5")
    assert result.sheet_name == "表5-1區域因素明細表(住)"
    assert result.source_digest == mapping.template_digest
    assert set(result.written_cells) == {"B2", "C5", "F5", "G5"}


def test_percentage_conventions_points_stay_points_and_fractions_scale() -> None:
    table5 = fill_workbook(
        _template("表5"),
        _mapping("table5-v1.json"),
        _snapshot({"table_5.P002.urban_plan_adjustment_pct": _present(Decimal("5"), unit=PCT)}),
    )
    assert "<v>5.00</v>" in _cell(_sheet_xml(table5, _template("表5")), "G5")

    table4 = fill_workbook(
        _template("表4"),
        _mapping("table4-v1.json"),
        _snapshot(
            {
                "table_4.P002.date_adjustment_pct": _present(Decimal("5"), unit=PCT),
                "table_4.P002.depth_adjustment_pct": _present(Decimal("2"), unit=PCT),
            }
        ),
    )
    xml = _sheet_xml(table4, _template("表4"))
    assert "<v>0.0500</v>" in _cell(xml, "J6")
    assert "<v>0.0200</v>" in _cell(xml, "J11")


def test_decimals_survive_without_a_float_round_trip() -> None:
    template = _template("表4")
    result = fill_workbook(
        template,
        _mapping("table4-v1.json"),
        _snapshot(
            {
                "table_4.P001.area": _present(Decimal("113.21"), unit="m2"),
                "table_4.P002.normal_unit_price": _present(Decimal("184763"), unit="TWD_per_m2"),
            }
        ),
    )
    xml = _sheet_xml(result, template)
    assert "<v>113.21</v>" in _cell(xml, "D9")
    assert "<v>184763</v>" in _cell(xml, "G5")


def test_missing_stays_blank_and_is_never_zero() -> None:
    template = _template("表5")
    result = fill_workbook(
        template,
        _mapping("table5-v1.json"),
        _snapshot(
            {
                "table_5.P002.urban_plan_grade": SnapshotEntry(state="missing"),
                "table_5.P003.urban_plan_adjustment_pct": SnapshotEntry(state="confirmed_zero"),
            },
            gaps={"table_5.P004.zone_number": "no third comparable this round"},
        ),
    )
    xml = _sheet_xml(result, template)
    assert _cell(xml, "E5").endswith("/>")  # missing: value removed, cell kept
    assert "<v>" not in _cell(xml, "E5")
    assert "<v>0.00</v>" in _cell(xml, "J5")  # confirmed zero is the only zero
    assert result.skipped["E5"] == "missing: snapshot marks the value missing"
    assert result.skipped["K3"] == "missing: no third comparable this round"
    assert "J5" in result.written_cells and "E5" not in result.written_cells


def test_wrong_template_bytes_are_refused_before_any_write() -> None:
    with pytest.raises(WorkbookWriteError) as failure:
        fill_workbook(
            _template("表5"),
            _mapping("table4-v1.json"),
            _snapshot(
                {
                    "table_4.case.case_number": _present("x"),
                }
            ),
        )
    assert failure.value.code == "template_digest_mismatch"


def test_text_that_looks_like_a_formula_stays_literal_text() -> None:
    template = _template("表5")
    result = fill_workbook(
        template,
        _mapping("table5-v1.json"),
        _snapshot({"table_5.P002.zone_number": _present("=SUM(A1)")}),
    )
    cell = _cell(_sheet_xml(result, template), "E3")
    assert 't="inlineStr"' in cell
    assert "<f>" not in cell
    assert "=SUM(A1)</t>" in cell

    import warnings

    from openpyxl import load_workbook

    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        workbook = load_workbook(io.BytesIO(result.content))
    assert workbook["表5-1區域因素明細表(住)"]["E3"].value == "=SUM(A1)"


def test_a_unit_that_contradicts_the_binding_is_a_hard_error() -> None:
    with pytest.raises(WorkbookWriteError) as failure:
        fill_workbook(
            _template("表5"),
            _mapping("table5-v1.json"),
            _snapshot({"table_5.P002.urban_plan_adjustment_pct": _present(Decimal("5"), unit="m")}),
        )
    assert failure.value.code == "unit_mismatch"


def test_sheet_count_and_visibility_are_unchanged() -> None:
    import warnings

    from openpyxl import load_workbook

    template = _template("表3")
    result = fill_workbook(
        template,
        _mapping("table3-v1.json"),
        _snapshot(
            {
                "table_3.case.zone_number": _present("P002-00"),
                "table_3.case.survey_date": _present("114年09月18日"),
                "table_3.case.main_road_width_m": _present(Decimal("18"), unit="m"),
            }
        ),
    )
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        before = load_workbook(io.BytesIO(template))
        after = load_workbook(io.BytesIO(result.content))
    assert after.sheetnames == before.sheetnames
    assert [after[n].sheet_state for n in after.sheetnames] == [
        before[n].sheet_state for n in before.sheetnames
    ]
    assert sum(1 for n in after.sheetnames if after[n].sheet_state == "visible") == 1


def test_not_applicable_renders_the_forms_own_dash_for_coverage_ratio() -> None:
    template = _template("表3")
    result = fill_workbook(
        template,
        _mapping("table3-v1.json"),
        _snapshot(
            {
                "table_3.case.building_coverage_text": SnapshotEntry(state="not_applicable"),
                "table_3.case.floor_area_ratio_text": _present("240%"),
            }
        ),
    )
    xml = _sheet_xml(result, template)
    assert ">-</t></is></c>" in _cell(xml, "H6")
    assert ">240%</t></is></c>" in _cell(xml, "H7")
