"""Build the real Shulin CalculationSnapshot from extracted + computed case data.

Bridges the two vocabularies this round produced:

- ``artifacts/shulin-case/{extraction,criteria,computed}.json`` (case-data agent)
- ``configs/mappings/table{3,4,5}-v1.json`` binding sources (official-writer agent)

Alignment is conservative and auditable: a key is matched either by
(subject segment, normalized Chinese label) or through the explicit alias
tables below; anything uncertain stays a gap with a written reason. Blank is
never zero, no equal-weight assumption is made, and the organizer's pre-filled
sample answers never enter the snapshot.

Percentage discipline: all percentage values are stored as percentage POINTS
(Decimal strings) with unit ``percent_points`` - the exact unit string the
mappings bind, so the writer's unit check passes. Table-4's fraction form is a
rendering decision the mapping makes (``percentage_fraction`` scales points by
10^-2 at write time); the snapshot never pre-divides.

End-to-end proof: after writing snapshot.json the script runs the real
``fill_workbook`` against the real templates in ``artifacts/official-templates``
and reports written/skipped cells plus sample values re-read via openpyxl.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import unicodedata
import warnings
from collections import Counter
from datetime import date
from decimal import Decimal
from pathlib import Path

from appraisal_review.adapters.local.workbook_writer import WorkbookWriteError, fill_workbook
from appraisal_review.domain.calculation_snapshot import (
    CalculationSnapshot,
    SnapshotEntry,
    SnapshotSubject,
)
from appraisal_review.domain.official_table_mapping import TableMapping
from appraisal_review.domain.service_contracts import RevisionReference

REPO = Path(__file__).resolve().parent.parent
ARTIFACTS = REPO / "artifacts" / "shulin-case"
MAPPINGS = REPO / "configs" / "mappings"
TEMPLATES = REPO / "artifacts" / "official-templates"

PERCENT_UNIT = "percent_points"  # the unit string every percentage binding declares
COMPARABLES = ("P002", "P003", "P004")
SUBJECT_PREFIXES = ("比較標的1", "比較標的2", "比較標的3", "比準地")
FACTOR_SUFFIXES = ("修正百分比", "優劣等級", "優劣註記", "差異率")

# --- explicit, hand-curated alias tables (audited against zh labels below) ---

# regional computed.json factor_id -> table_5 field stem
REGIONAL_STEM = {
    "urban_plan_status": "urban_plan",
    "zoning_district": "zoning",
    "coverage_ratio": "building_coverage",
    "floor_area_ratio": "floor_area_ratio",
    "building_prohibition": "building_prohibition",
    "building_restriction": "building_restriction",
    "main_road_width": "main_road_width",
    "avg_road_width_in_section": "avg_road_width",
    "large_station_proximity": "major_station_access",
    "bus_stop_proximity": "bus_stop_access",
    "interchange_proximity": "interchange_access",
    "road_planning_development": "road_development",
    "sunlight": "sunlight",
    "landscape": "landscape",
    "slope": "slope",
    "drainage": "drainage",
    "terrain": "terrain",
    "site_improvement": "land_improvement",
    "school_proximity": "school_access",
    "market_proximity": "market_access",
    "park_plaza_proximity": "park_access",
    "tourism_recreation_proximity": "tourism_access",
    "parking_availability": "parking_access",
    "service_facility_proximity": "service_facility_access",
    "utility_gas_facility": "utility_facility",
    "funeral_facility": "funeral_facility",
    "waste_facility": "waste_facility",
    "environmental_pollution": "pollution",
    "other_factors": "other_factor",
}

# individual computed.json factor_id -> table_4 field stem
INDIVIDUAL_STEM = {
    "parcel_area": "area",
    "parcel_width": "width",
    "parcel_depth": "depth",
    "parcel_shape": "shape",
    "street_frontage": "frontage",
    "parcel_terrain": "terrain",
    "road_type": "road_type",
    "frontage_road_width": "frontage_road",
    "school_proximity_parcel": "school",
    "market_proximity_parcel": "market",
    "park_plaza_proximity_parcel": "park",
    "station_proximity_parcel": "station",
    "commercial_district_proximity": "business_district",
    "nuisance_facility": "nuisance_facility",
    "parking_convenience": "parking",
    "zoning_district_parcel": "zoning",
    "coverage_ratio_parcel": "building_coverage",
    "far_parcel": "floor_area_ratio",
    "no_build_restriction_parcel": "building_restriction",
    "dead_end_alley": "other",
}

# alias pairs whose normalized zh labels legitimately differ (wording variants
# between the criteria sheet's row title and the workbook's shorter row title);
# every other alias must agree label-for-label or it is demoted to a gap.
APPROVED_LABEL_VARIANTS = {
    ("土地改良", "建築基地改良或其他改良"),
    ("交流道之有無及接近程度", "交流道之有無及接近交流道之程度"),
    ("接近公園、廣場之程度", "接近公園、廣場、徒步區之程度"),
    ("電業設施及公用氣體燃料設施", "電業設施及公用氣體燃料設施之有無及接近程度"),
    ("殯葬設施", "殯葬設施之有無及接近程度"),
    ("廢棄物處理設施", "廢棄物處理設施之有無及接近程度"),
    ("環境污染", "水污染、噪音污染、廢氣污染、廢棄物污染等之有無及接近程度"),
    ("使用分區或編定用地", "使用分區或編定"),
    ("排水之良否", "保排水之良否"),
}

GAP_NO_ALIGNMENT = "no confident alignment from extracted data"
GAP_GRADE_CODE = (
    "no confident alignment from extracted data: the sheet's numeric grade-code "
    "convention (优/劣 word to 1..N code) was not established from the assignment"
)
GAP_DOWNSTREAM = (
    "downstream of missing/blocked inputs (parcel attributes, zoning confirmation, "
    "FAR method); multi-comparable weight rule taken from the manual only, never "
    "assumed equal - not computed in this snapshot"
)


def normalize_label(text: str) -> str:
    """NFKC-fold, drop parenthetical qualifiers and separators for label matching."""

    text = unicodedata.normalize("NFKC", text)
    text = text.replace("（", "(").replace("）", ")").replace("，", "、")
    text = re.sub(r"\([^()]*\)", "", text)
    text = re.sub(r"\([^()]*\)", "", text)  # once more for formerly nested groups
    text = re.sub(r"[\s/·・…]+", "", text)
    return text.strip()


def factor_label(binding_label: str) -> str:
    """Strip subject prefixes and role suffixes from a table_4/5 binding label."""

    text = normalize_label(binding_label)
    for prefix in SUBJECT_PREFIXES:
        if text.startswith(prefix):
            text = text[len(prefix) :]
            break
    for suffix in FACTOR_SUFFIXES:
        if text.endswith(suffix):
            text = text[: -len(suffix)]
            break
    return text


def clip_trace(text: str) -> str:
    return text if len(text) <= 2048 else text[:2045] + "..."


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def present(
    value: Decimal | str,
    origin: str,
    trace: str,
    unit: str | None = None,
) -> SnapshotEntry:
    return SnapshotEntry(
        state="present", value=value, unit=unit, origin=origin, trace=clip_trace(trace)
    )


def not_applicable(trace: str) -> SnapshotEntry:
    return SnapshotEntry(state="not_applicable", trace=clip_trace(trace))


class Builder:
    def __init__(self) -> None:
        self.mappings: dict[int, TableMapping] = {}
        self.binding_by_source: dict[str, tuple[int, object]] = {}
        for table in (3, 4, 5):
            mapping = TableMapping.model_validate(load_json(MAPPINGS / f"table{table}-v1.json"))
            self.mappings[table] = mapping
            for binding in mapping.bindings:
                self.binding_by_source[binding.source] = (table, binding)
        self.demand = set(self.binding_by_source)
        self.extraction = load_json(ARTIFACTS / "extraction.json")
        self.computed = load_json(ARTIFACTS / "computed.json")
        self.criteria = load_json(ARTIFACTS / "criteria.json")
        self.entries: dict[str, SnapshotEntry] = {}
        self.gaps: dict[str, str] = {}
        self.alias_audit: list[tuple[str, str, str, str]] = []  # alias, binding zh, data zh, verdict
        self.blank_facility_reason = next(
            r["reason"]
            for r in self.computed["regional_corrections"]
            if r["state"] == "missing"
        )

    # -- helpers -------------------------------------------------------------

    def put(self, key: str, entry: SnapshotEntry) -> None:
        if key not in self.demand:
            raise KeyError(f"snapshot tried to provide a key no mapping binds: {key}")
        if key in self.gaps:
            return  # an explicit gap (e.g. awaiting confirmation) wins over a value
        self.entries[key] = entry

    def gap(self, key: str, reason: str) -> None:
        if key in self.demand and key not in self.entries:
            self.gaps[key] = reason

    def alias_checked(self, stem_key: str, data_zh: str, alias_name: str) -> bool:
        """Verify an alias by comparing normalized labels; demote to gap on mismatch."""

        found = self.binding_by_source.get(stem_key)
        if found is None:
            self.alias_audit.append((alias_name, "<no binding>", data_zh, "missing-binding"))
            return False
        _, binding = found
        binding_side = factor_label(binding.label)
        data_side = normalize_label(data_zh)
        ok = binding_side == data_side or (binding_side, data_side) in APPROVED_LABEL_VARIANTS
        self.alias_audit.append(
            (alias_name, binding_side, data_side, "ok" if ok else "MISMATCH->gap")
        )
        return ok

    # -- explicit awaiting-confirmation gaps first (they must win) ------------

    def block_open_conditions(self) -> None:
        zoning = next(
            c
            for c in self.extraction["open_conditions"]
            if c["id"] == "p001_zoning_mrt_development_zone"
        )
        detail = (
            "P001 使用分區 field says 第一種住宅區 but the 區段範圍 text reads 捷運開發區"
            "(變更前為第一種住宅區); grades differ by one step (+3.75 points per comparable "
            "in 表4 row 22 if 優 is confirmed). " + zoning["note_on_wording"]
        )
        reason = f"awaiting human confirmation: {detail}"
        for key in ("table_3.case.zoning", "table_4.P001.zoning"):
            self.gap(key, reason)
        for comp in COMPARABLES:
            self.gap(f"table_4.{comp}.zoning_adjustment_pct", reason)
        # table_3 grade cells for the ambiguous zoning row
        self.gap("table_3.case.zoning_grade", reason)
        self.gap("table_3.case.zoning_grade_scale", reason)

    # -- structural / given-input entries -------------------------------------

    def add_case_fields(self) -> None:
        case = self.extraction["case"]
        number = case["case_number"]
        for key in ("table_4.case.case_number", "table_5.case.case_number"):
            self.put(
                key,
                present(
                    number["value"],
                    "given_input",
                    f"assignment p.{number['page']} quote: {number['quote']}",
                ),
            )
        vdate = case["valuation_date"]
        self.put(
            "table_4.case.valuation_date",
            present(
                vdate["roc"],
                "given_input",
                f"assignment p.{vdate['page']} quote: {vdate['quote']} (ROC 1110901 = 2022-09-01)",
            ),
        )
        serial = case["base_parcel_serial"]
        self.put(
            "table_4.P001.serial_number",
            present(
                serial["value"],
                "given_input",
                f"assignment p.{serial['page']} ({serial['zh_label']}) quote: {serial['quote']}",
            ),
        )
        pdn = case["price_date_note"]
        ccn = case["case_collection_note"]
        self.put(
            "table_4.case.note",
            present(
                f"{pdn['quote']}\n{ccn['quote']}",
                "given_input",
                f"verbatim 備註 quotes from assignment p.{pdn['page']} and p.{ccn['page']}",
            ),
        )
        remark = self.extraction["table5_1"]["regional_move_remark"]
        self.put(
            "table_5.case.note",
            present(
                remark["quote"],
                "given_input",
                f"表5-1 remark, assignment p.{remark['page']} quote: {remark['quote']}",
            ),
        )

    def add_subject_structurals(self) -> None:
        subjects = self.extraction["subjects"]
        for sid, sub in subjects.items():
            code = sub["attributes"]["section_code"]
            trace = f"assignment p.{code['page']} ({code['zh_label']}) value {code['value']}"
            for key in (f"table_4.{sid}.zone_number", f"table_5.{sid}.zone_number"):
                if key in self.demand:
                    self.put(key, present(code["value"], "given_input", trace))
        table4 = self.extraction["table4"]
        p1 = table4["P001"]["address"]
        self.put(
            "table_4.P001.land_description",
            present(
                p1["value"], "given_input", f"assignment p.{p1['page']} quote: {p1['quote']}"
            ),
        )
        for comp in COMPARABLES:
            row = table4[comp]
            addr = row["address"]
            self.put(
                f"table_4.{comp}.land_description",
                present(
                    addr["value"],
                    "given_input",
                    f"assignment p.{addr['page']} quote: {addr['quote']}",
                ),
            )
            tdate = row["transaction_date"]
            self.put(
                f"table_4.{comp}.transaction_date",
                present(
                    tdate["value"],
                    "given_input",
                    f"assignment p.{tdate['page']} quote: {tdate['quote']}",
                ),
            )
            price = row["normal_unit_price"]
            self.put(
                f"table_4.{comp}.normal_unit_price",
                present(
                    Decimal(price["value"]),
                    "given_input",
                    f"assignment p.{price['page']} quote: {price['quote']} (TWD/m2)",
                    unit="TWD_per_m2",
                ),
            )
            rate = row["price_date_adjustment_rate"]
            self.put(
                f"table_4.{comp}.date_adjustment_pct",
                present(
                    Decimal(rate["value"]),
                    "given_input",
                    f"assignment p.{rate['page']} quote: {rate['quote']} - stored as "
                    "percentage points; the table_4 mapping renders the fraction form",
                    unit=PERCENT_UNIT,
                ),
            )
            adjusted = row["adjusted_unit_price"]
            check = next(
                c
                for c in self.computed["price_date_adjustment_checks"]
                if c["comparable"] == comp
            )
            self.put(
                f"table_4.{comp}.adjusted_unit_price",
                present(
                    Decimal(adjusted["value"]),
                    "given_input",
                    f"assignment p.{adjusted['page']} quote: {adjusted['quote']}; "
                    f"independently verified: {check['trace']}",
                    unit="TWD_per_m2",
                ),
            )

    def add_table4_conditions(self) -> None:
        """表4 rows 22-25 input values that come straight from the 表3 sheets."""

        subjects = self.extraction["subjects"]
        for sid, sub in subjects.items():
            attrs = sub["attributes"]
            coverage = attrs["coverage_ratio"]
            key = f"table_4.{sid}.building_coverage"
            if key in self.demand:
                self.put(
                    key,
                    present(
                        Decimal(coverage["value"]),
                        "given_input",
                        f"assignment p.{coverage['page']} 建蔽率 {coverage['value']}% - "
                        "points stored; mapping renders the fraction",
                        unit=PERCENT_UNIT,
                    ),
                )
            far = attrs["floor_area_ratio"]
            key = f"table_4.{sid}.floor_area_ratio"
            if key in self.demand:
                self.put(
                    key,
                    present(
                        Decimal(far["value"]),
                        "given_input",
                        f"assignment p.{far['page']} 容積率 {far['value']}% - points stored; "
                        "mapping renders the fraction",
                        unit=PERCENT_UNIT,
                    ),
                )
            prohibition = attrs["building_prohibition"]
            restriction = attrs["building_restriction"]
            key = f"table_4.{sid}.building_restriction"
            if (
                key in self.demand
                and prohibition["value"] == "無"
                and restriction["value"] == "無"
            ):
                self.put(
                    key,
                    present(
                        "無",
                        "given_input",
                        f"assignment p.{prohibition['page']}: 有無禁止建築=無 and "
                        "有無限制建築=無, so 禁限建=無",
                    ),
                )
            zoning = attrs["zoning_district"]
            key = f"table_4.{sid}.zoning"
            if key in self.demand and key not in self.gaps:  # P001 already gapped
                self.put(
                    key,
                    present(
                        zoning["value"],
                        "given_input",
                        f"assignment p.{zoning['page']} 使用分區 {zoning['value']}",
                    ),
                )

    # -- table_3: comparison base zone sheet, matched by normalized zh label ---

    def add_table3(self) -> None:
        attrs = self.extraction["subjects"]["P001"]["attributes"]
        index: dict[str, tuple[str, dict]] = {}
        for name, attr in attrs.items():
            if not isinstance(attr, dict) or "zh_label" not in attr:
                continue
            if isinstance(attr.get("value"), list):
                continue  # checkbox lists are not cell-mappable
            index[normalize_label(attr["zh_label"])] = (name, attr)
        percent_text = {"building_coverage_text", "floor_area_ratio_text"}
        for binding in self.mappings[3].bindings:
            source = binding.source
            field = source.split(".")[-1]
            if source in self.gaps or field.endswith(("_grade", "_grade_scale")):
                continue
            matched = index.get(normalize_label(binding.label))
            if matched is None:
                continue
            name, attr = matched
            if name == "zoning_district":
                continue  # explicitly blocked by the open condition
            value = attr["value"]
            trace = (
                f"assignment p.{attr['page']} ({attr['zh_label']}) value {value}"
                + (f" {attr['unit']}" if attr.get("unit") else "")
            )
            if binding.value_kind == "decimal":
                self.put(
                    source,
                    present(Decimal(value), "given_input", trace, unit=binding.unit),
                )
            elif binding.value_kind in {"text", "date"}:
                text = value
                if field in percent_text and attr.get("unit") == "percent":
                    text = f"{value}%"
                self.put(source, present(text, "given_input", trace))

    # -- computed corrections --------------------------------------------------

    def add_regional(self) -> None:
        rows = self.computed["regional_corrections"]
        p001_grade_seen: dict[str, tuple[str, str]] = {}
        for row in rows:
            fid = row["factor_id"]
            stem = REGIONAL_STEM.get(fid)
            comp = row["comparable"]
            if stem is None:
                continue
            adj_key = f"table_5.{comp}.{stem}_adjustment_pct"
            if not self.alias_checked(adj_key, row["zh_label"], f"regional:{fid}->{stem}"):
                self.gap(adj_key, GAP_NO_ALIGNMENT)
                continue
            grade_key = f"table_5.{comp}.{stem}_grade"
            label_key = f"table_5.{comp}.{stem}_grade_label"
            state = row["state"]
            if state == "computed":
                self.put(
                    adj_key,
                    present(
                        Decimal(row["correction_points"]),
                        "computed",
                        row["trace"],
                        unit=PERCENT_UNIT,
                    ),
                )
                self.put(label_key, present(row["comparable_grade"], "computed", row["trace"]))
                self.gap(grade_key, GAP_GRADE_CODE)
                p001_grade_seen[stem] = (row["p001_grade"], row["trace"])
            elif state == "given":
                trace = row["source"]
                self.put(
                    adj_key,
                    present(Decimal(row["correction_points"]), "given_input", trace, unit=PERCENT_UNIT),
                )
                self.put(label_key, present("無", "given_input", trace))
                self.gap(grade_key, GAP_GRADE_CODE)
            elif state == "excluded_by_assignment_remark":
                for key in (adj_key, grade_key, label_key):
                    if key in self.demand:
                        self.put(key, not_applicable(row["reason"]))
            elif state == "missing":
                for key in (adj_key, grade_key, label_key):
                    self.gap(key, row["reason"])
        # P001 grade labels (grades are shared across the three comparable rows)
        for stem, (grade, trace) in p001_grade_seen.items():
            label_key = f"table_5.P001.{stem}_grade_label"
            if label_key in self.demand:
                self.put(label_key, present(grade, "computed", trace))
            self.gap(f"table_5.P001.{stem}_grade", GAP_GRADE_CODE)
        # P001 side of excluded / missing / given rows
        for row in rows:
            if row["comparable"] != "P002":
                continue
            stem = REGIONAL_STEM.get(row["factor_id"])
            if stem is None:
                continue
            for key in (f"table_5.P001.{stem}_grade", f"table_5.P001.{stem}_grade_label"):
                if key not in self.demand or key in self.entries or key in self.gaps:
                    continue
                if row["state"] == "excluded_by_assignment_remark":
                    self.put(key, not_applicable(row["reason"]))
                elif row["state"] == "missing":
                    self.gap(key, row["reason"])
                elif row["state"] == "given":
                    if key.endswith("_grade_label"):
                        self.put(key, present("無", "given_input", row["source"]))
                    else:
                        self.gap(key, GAP_GRADE_CODE)

    def add_regional_subtotals(self) -> None:
        for record in self.computed["regional_group_subtotals"]:
            match = re.search(r"\((\d)\)", record["group"])
            if match is None:
                continue
            number = match.group(1)
            for comp in COMPARABLES:
                key = f"table_5.{comp}.group{number}_subtotal_pct"
                if record["status"] in {"complete", "given"} and record.get("subtotal_points"):
                    value = record["subtotal_points"][comp]
                    trace = (
                        f"sum of {record['group']} computed rows for {comp} = {value} points"
                    )
                    if record["status"] == "given":
                        trace = (
                            f"{record['group']} holds a single assignment-given row "
                            f"(其他影響因素 0.00, p.5); subtotal = {value} points"
                        )
                    if record.get("excluded_factors_moved_to_table4"):
                        trace += (
                            "; excludes "
                            + "/".join(record["excluded_factors_moved_to_table4"])
                            + " per the 表5-1 remark (corrected in 表4)"
                        )
                    self.put(key, present(Decimal(value), "computed", trace, unit=PERCENT_UNIT))
                else:
                    missing = "、".join(record.get("missing_factors", []))
                    self.gap(
                        key,
                        f"{record['group']} subtotal not computable: factors missing "
                        f"inputs ({missing}); a partial subtotal must not pose as complete",
                    )
        caveat = self.computed["regional_partial_totals"]["caveat"]
        for comp in COMPARABLES:
            self.gap(f"table_5.{comp}.total_adjustment_pct", caveat)
            self.gap(f"table_4.{comp}.region_adjustment_pct", caveat)

    def add_individual(self) -> None:
        missing_stem_reasons: dict[str, str] = {}
        for row in self.computed["individual_corrections"]:
            stem = INDIVIDUAL_STEM.get(row["factor_id"])
            comp = row["comparable"]
            if stem is None:
                continue
            adj_key = f"table_4.{comp}.{stem}_adjustment_pct"
            state = row["state"]
            if state == "computed":
                if not self.alias_checked(
                    adj_key, row["zh_label"], f"individual:{row['factor_id']}->{stem}"
                ):
                    self.gap(adj_key, GAP_NO_ALIGNMENT)
                    continue
                self.put(
                    adj_key,
                    present(
                        Decimal(row["correction_points"]),
                        "computed",
                        row["trace"],
                        unit=PERCENT_UNIT,
                    ),
                )
            elif state in {"blocked_by_open_condition", "manual_method_required"}:
                self.gap(adj_key, f"awaiting human confirmation: {row['reason']}")
            elif state == "missing":
                missing_stem_reasons[stem] = row["reason"]
                self.gap(adj_key, row["reason"])
        # the blank parcel-attribute value cells share the factor rows' reasons
        for source in self.demand:
            if not source.startswith("table_4."):
                continue
            field = source.split(".")[-1]
            if source in self.entries or source in self.gaps:
                continue
            for stem, reason in missing_stem_reasons.items():
                if field == stem or field.startswith(f"{stem}_"):
                    self.gap(source, reason)
                    break

    # -- final sweep: every remaining demanded key becomes an honest gap -------

    def sweep(self) -> None:
        downstream_fields = {
            "total_adjustment_pct",
            "abs_adjustment_sum_pct",
            "similarity",
            "trial_price",
            "weight_pct",
            "compared_price",
        }
        for source in sorted(self.demand):
            if source in self.entries or source in self.gaps:
                continue
            table, _, field = source.split(".", 2)
            field = field.split(".")[-1]
            if table == "table_3":
                if field.endswith(("_grade", "_grade_scale")):
                    self.gap(source, GAP_GRADE_CODE)
                elif field.endswith(("_name", "_distance_m", "_count")):
                    self.gap(source, self.blank_facility_reason)
                else:
                    self.gap(source, GAP_NO_ALIGNMENT)
            elif field in downstream_fields:
                self.gap(source, GAP_DOWNSTREAM)
            else:
                self.gap(source, GAP_NO_ALIGNMENT)

    # -- assembly ---------------------------------------------------------------

    def build(self, case_id: str, revision_id: str, material_digest: str) -> CalculationSnapshot:
        self.block_open_conditions()
        self.add_case_fields()
        self.add_subject_structurals()
        self.add_table4_conditions()
        self.add_table3()
        self.add_regional()
        self.add_regional_subtotals()
        self.add_individual()
        self.sweep()
        return CalculationSnapshot(
            revision=RevisionReference(
                case_id=case_id, revision_id=revision_id, material_digest=material_digest
            ),
            district="Shulin",
            valuation_date=date(2022, 9, 1),
            rule_bundle_id="shulin-1110901",
            rule_bundle_version="v1",
            subjects=(
                SnapshotSubject(subject_id="P001", role="comparison_base", label="比準地"),
                SnapshotSubject(subject_id="P002", role="comparable", label="比較標的1"),
                SnapshotSubject(subject_id="P003", role="comparable", label="比較標的2"),
                SnapshotSubject(subject_id="P004", role="comparable", label="比較標的3"),
            ),
            entries=self.entries,
            gaps=self.gaps,
        )


def default_material_digest() -> str:
    digest = hashlib.sha256()
    for name in ("extraction.json", "criteria.json", "computed.json"):
        digest.update((ARTIFACTS / name).read_bytes())
    return digest.hexdigest()


def find_template(mapping: TableMapping) -> Path | None:
    for path in sorted(TEMPLATES.glob("*.xlsx")):
        if hashlib.sha256(path.read_bytes()).hexdigest() == mapping.template_digest:
            return path
    return None


SAMPLE_SOURCES = {
    3: [
        "table_3.case.year_term",
        "table_3.case.zone_number",
        "table_3.case.zone_extent",
        "table_3.case.urban_plan",
        "table_3.case.main_road_width_m",
        "table_3.case.avg_road_width_m",
        "table_3.case.building_coverage_text",
        "table_3.case.drainage",
    ],
    4: [
        "table_4.case.valuation_date",
        "table_4.case.case_number",
        "table_4.P001.serial_number",
        "table_4.P002.normal_unit_price",
        "table_4.P002.date_adjustment_pct",
        "table_4.P002.building_coverage",
        "table_4.P002.building_restriction_adjustment_pct",  # computed correction
        "table_4.P002.zone_number",
    ],
    5: [
        "table_5.case.case_number",
        "table_5.P001.zone_number",
        "table_5.P002.sunlight_adjustment_pct",  # computed correction
        "table_5.P002.urban_plan_grade_label",
        "table_5.P002.other_factor_adjustment_pct",  # assignment-given 0
        "table_5.P002.group1_subtotal_pct",
        "table_5.P003.avg_road_width_adjustment_pct",  # computed correction
        "table_5.P002.group3_subtotal_pct",
    ],
}


def run_writer_proof(builder: Builder, snapshot: CalculationSnapshot, out_dir: Path) -> list[str]:
    from openpyxl import load_workbook

    out_dir.mkdir(parents=True, exist_ok=True)
    lines: list[str] = []
    for table, mapping in sorted(builder.mappings.items()):
        template = find_template(mapping)
        if template is None:
            lines.append(
                f"table {table}: NO template in {TEMPLATES} matches digest "
                f"{mapping.template_digest} - writer not run"
            )
            continue
        try:
            filled = fill_workbook(template.read_bytes(), mapping, snapshot)
        except WorkbookWriteError as error:
            lines.append(f"table {table}: writer refused: {error.code}: {error.detail}")
            continue
        out_path = out_dir / f"table{table}-filled.xlsx"
        out_path.write_bytes(filled.content)
        reasons = Counter(reason.split(":")[0] + ": " + reason.split(":", 1)[1][:60]
                          for reason in filled.skipped.values())
        lines.append(
            f"table {table}: template {template.name}\n"
            f"  written cells: {len(filled.written_cells)}, skipped: {len(filled.skipped)}"
        )
        for reason, count in reasons.most_common(4):
            lines.append(f"  skip reason x{count}: {reason}")
        lines.append(f"  output: {out_path}")
        lines.append(f"  sha256: {hashlib.sha256(filled.content).hexdigest()}")
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            workbook = load_workbook(io.BytesIO(filled.content))
        sheet = workbook[mapping.sheet_name]
        cell_of = {b.source: b.cell for b in mapping.bindings}
        for source in SAMPLE_SOURCES[table]:
            cell = cell_of.get(source)
            if cell is None:
                lines.append(f"  sample {source}: (not bound)")
                continue
            lines.append(f"  sample {source} [{cell}] = {sheet[cell].value!r}")
    return lines


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--case-id", default="shulin-1110901-99-XXX")
    parser.add_argument("--revision-id", default="rev-local-1")
    parser.add_argument("--material-digest", default=None)
    parser.add_argument("--output", default=str(ARTIFACTS / "snapshot.json"))
    args = parser.parse_args()

    builder = Builder()
    snapshot = builder.build(
        args.case_id, args.revision_id, args.material_digest or default_material_digest()
    )
    snapshot = CalculationSnapshot.model_validate(snapshot.model_dump(mode="json"))

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(snapshot.model_dump_json(indent=2), encoding="utf-8")

    per_table: dict[int, Counter] = {3: Counter(), 4: Counter(), 5: Counter()}
    for source in builder.demand:
        table = int(source.split(".")[0].split("_")[1])
        if source in snapshot.entries:
            state = snapshot.entries[source].state
            per_table[table]["provided" if state == "present" else state] += 1
        else:
            per_table[table]["gaps"] += 1

    print(f"snapshot written: {output}")
    print(f"snapshot digest: {snapshot.digest()}")
    print(f"entries: {len(snapshot.entries)} (gaps: {len(snapshot.gaps)}; "
          f"demand {len(builder.demand)})")
    for table in (3, 4, 5):
        counts = per_table[table]
        bound = sum(counts.values())
        print(
            f"  table {table}: bound {bound}, present {counts['provided']}, "
            f"not_applicable {counts['not_applicable']}, gaps {counts['gaps']}"
        )
    mismatches = [row for row in builder.alias_audit if row[3] != "ok"]
    if mismatches:
        print("alias audit mismatches (demoted to gaps):")
        for row in mismatches:
            print("  ", row)

    print("\n-- writer proof (real templates, real fill_workbook) --")
    for line in run_writer_proof(builder, snapshot, ARTIFACTS / "filled"):
        print(line)

    important_gaps = [
        "table_4.P001.compared_price",
        "table_4.P002.weight_pct",
        "table_5.P002.total_adjustment_pct",
        "table_4.P002.region_adjustment_pct",
        "table_4.P002.zoning_adjustment_pct",
    ]
    print("\n5 most important unmatched keys:")
    for key in important_gaps:
        print(f"  {key}: {snapshot.gaps.get(key, '(provided)')[:110]}")


if __name__ == "__main__":
    main()
